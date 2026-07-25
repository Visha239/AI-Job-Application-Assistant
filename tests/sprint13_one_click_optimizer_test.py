from __future__ import annotations

from pathlib import Path
import tempfile

from docx import Document

from app.services.one_click_optimizer import (
    prepare_one_click_application,
)


def build_test_resume(
    path: Path,
) -> None:
    document = Document()
    document.add_paragraph(
        "VISHAL BANAKAR"
    )
    document.add_heading(
        "CAREER OBJECTIVE",
        level=1,
    )
    document.add_paragraph(
        "Looking for an analytics role."
    )
    document.add_heading(
        "TECHNICAL SKILLS",
        level=1,
    )
    document.add_paragraph(
        "Excel, SQL, Python, Power BI, Linux"
    )
    document.add_heading(
        "WORK EXPERIENCE",
        level=1,
    )
    document.add_paragraph(
        "Worked on SQL reports and production support."
    )
    document.add_heading(
        "PROJECTS",
        level=1,
    )
    document.add_paragraph(
        "Made a Power BI dashboard."
    )
    document.add_heading(
        "EDUCATION",
        level=1,
    )
    document.add_paragraph(
        "B.Tech Computer Science"
    )
    document.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        resume_path = (
            Path(temp_dir)
            / "resume.docx"
        )

        build_test_resume(
            resume_path
        )

        session_state: dict = {}

        selected_job = {
            "company": (
                "CareerPilot Test Company"
            ),
            "job_role": "Data Analyst",
            "location": "Bengaluru",
            "job_link": (
                "https://example.com/job"
            ),
            "job_description": (
                "We need a Data Analyst with SQL, Python, "
                "Power BI, Excel, Dashboard, Reporting, "
                "Azure and Snowflake experience. "
                "The candidate will prepare reports, analyze data, "
                "build dashboards and work with stakeholders."
            ),
            "match_score": 88,
            "priority": "High priority",
            "apply_urgency": (
                "Apply today"
            ),
            "source": "linkedin",
        }

        result = (
            prepare_one_click_application(
                session_state=(
                    session_state
                ),
                profile={
                    "name": (
                        "Vishal Basavaraj Banakar"
                    ),
                    "email": (
                        "vishalbanakar321@gmail.com"
                    ),
                    "skills": [
                        "SQL",
                        "Python",
                        "Power BI",
                        "Excel",
                        "Dashboard",
                        "Reporting",
                    ],
                },
                selected_job=(
                    selected_job
                ),
                resume_type=(
                    "Data Analyst"
                ),
                source_resume_path=(
                    resume_path
                ),
                create_crm_record=False,
                create_outreach_draft=False,
            )
        )

        assert result[
            "ats_score"
        ] > 0
        assert "SQL" in result[
            "matched_skills"
        ]
        assert "Azure" in result[
            "missing_skills"
        ]
        assert Path(
            result[
                "final_resume_path"
            ]
        ).exists()
        assert Path(
            result[
                "summary_path"
            ]
        ).exists()
        assert session_state[
            "careerpilot_selected_job"
        ]["company"] == (
            "CareerPilot Test Company"
        )

        print("=" * 70)
        print(
            "CAREERPILOT SPRINT 13 "
            "ONE-CLICK OPTIMIZER TEST"
        )
        print("=" * 70)
        print(
            "ATS score:",
            result[
                "ats_score"
            ],
        )
        print(
            "Matched:",
            result[
                "matched_skills"
            ],
        )
        print(
            "Missing:",
            result[
                "missing_skills"
            ],
        )
        print(
            "Resume:",
            result[
                "final_resume_path"
            ],
        )
        print(
            "\nSPRINT 13 ONE-CLICK OPTIMIZER TEST PASSED"
        )


if __name__ == "__main__":
    main()
