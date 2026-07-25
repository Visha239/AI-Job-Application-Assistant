from __future__ import annotations

from pathlib import Path
import tempfile

from docx import Document

from app.services.resume_intelligence_v2 import (
    generate_tailored_resume,
    review_resume,
)


def build_test_resume(
    path: Path,
) -> None:
    document = Document()
    document.add_paragraph(
        "VISHAL BANAKAR"
    )
    document.add_heading(
        "CAREER OBJECTIVE",
        level=1,
    )
    document.add_paragraph(
        "Looking for an opportunity in analytics."
    )
    document.add_heading(
        "TECHNICAL SKILLS",
        level=1,
    )
    document.add_paragraph(
        "Excel, SQL, Python, Power BI, Linux"
    )
    document.add_heading(
        "WORK EXPERIENCE",
        level=1,
    )
    document.add_paragraph(
        "Worked on SQL reports and Linux troubleshooting."
    )
    document.add_paragraph(
        "Resolved production issues using ServiceNow and Jira."
    )
    document.add_heading(
        "PROJECTS",
        level=1,
    )
    document.add_paragraph(
        "Made a Power BI dashboard for manufacturing analysis."
    )
    document.add_heading(
        "EDUCATION",
        level=1,
    )
    document.add_paragraph(
        "B.Tech Computer Science"
    )
    document.save(path)


def main() -> None:
    with tempfile.TemporaryDirectory() as temp_dir:
        resume_path = (
            Path(temp_dir)
            / "resume.docx"
        )
        output_path = (
            Path(temp_dir)
            / "tailored.docx"
        )

        build_test_resume(
            resume_path
        )

        jd = """
        We need a Data Analyst with SQL, Python, Power BI,
        Excel, Dashboard, Reporting, Data Visualization,
        Stakeholder Management, Azure, and Snowflake.
        """

        review = review_resume(
            resume_path=resume_path,
            job_description=jd,
            profile_skills=[
                "SQL",
                "Python",
                "Power BI",
                "Excel",
                "Linux",
                "ServiceNow",
                "Jira",
            ],
        )

        assert "SQL" in review[
            "matched_skills"
        ]
        assert "Power BI" in review[
            "matched_skills"
        ]
        assert "Azure" in review[
            "missing_skills"
        ]
        assert "Snowflake" in review[
            "missing_skills"
        ]
        assert review[
            "bullet_reviews"
        ]
        assert review[
            "recommendations"
        ]

        result = generate_tailored_resume(
            source_resume_path=resume_path,
            output_path=output_path,
            job_description=jd,
            role="Data Analyst",
            profile_skills=[
                "SQL",
                "Python",
                "Power BI",
                "Excel",
                "Linux",
            ],
        )

        assert output_path.exists()
        assert result[
            "objective_updated"
        ] is True
        assert result[
            "skills_reordered"
        ] is True
        assert result[
            "bullets_improved"
        ] >= 1

        final_document = Document(
            output_path
        )
        final_text = "\n".join(
            paragraph.text
            for paragraph
            in final_document.paragraphs
        )

        assert (
            "Stakeholder Management"
            not in final_text
        )
        assert (
            "Snowflake"
            not in final_text
        )
        assert (
            "Azure"
            not in final_text
        )

        print("=" * 70)
        print(
            "CAREERPILOT SPRINT 11 "
            "RESUME INTELLIGENCE TEST"
        )
        print("=" * 70)
        print(
            "Overall score:",
            review[
                "overall_score"
            ],
        )
        print(
            "Matched skills:",
            review[
                "matched_skills"
            ],
        )
        print(
            "Missing skills:",
            review[
                "missing_skills"
            ],
        )
        print(
            "Bullets improved:",
            result[
                "bullets_improved"
            ],
        )
        print(
            "\nSPRINT 11 RESUME INTELLIGENCE TEST PASSED"
        )


if __name__ == "__main__":
    main()
