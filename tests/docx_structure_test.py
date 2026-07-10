from docx import Document

doc = Document("resumes/master/master_resume.docx")

print("=" * 60)
print("PARAGRAPHS")
print("=" * 60)

for i, p in enumerate(doc.paragraphs):
    print(i, ":", repr(p.text))

print("\n")
print("=" * 60)
print("TABLES")
print("=" * 60)

for t, table in enumerate(doc.tables):
    print(f"\nTable {t}")

    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            print(f"Row {r} Cell {c}")
            print(repr(cell.text))
            print("-" * 40)