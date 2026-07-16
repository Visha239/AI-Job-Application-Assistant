from __future__ import annotations

import re
from collections import Counter
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document


DEFAULT_MASTER_RESUME = Path(
    "resumes/master/master_resume.docx"
)
DEFAULT_OUTPUT_DIR = Path(
    "exports/docx/resume_intelligence"
)

SECTION_ALIASES = {
    "objective": {
        "career objective",
        "professional summary",
        "summary",
        "profile",
    },
    "skills": {
        "technical skills",
        "skills",
        "core skills",
        "key skills",
        "technical proficiency",
    },
    "experience": {
        "work experience",
        "professional experience",
        "experience",
        "employment history",
    },
    "projects": {
        "projects",
        "academic projects",
        "key projects",
        "project experience",
    },
    "education": {
        "education",
        "academic qualification",
        "academic qualifications",
    },
    "certifications": {
        "certifications",
        "certificates",
        "courses",
    },
}

KNOWN_SKILLS = [
    "SQL",
    "Python",
    "Power BI",
    "Excel",
    "Tableau",
    "DAX",
    "Power Query",
    "ETL",
    "Data Modeling",
    "Data Analysis",
    "Data Visualization",
    "Dashboard",
    "Reporting",
    "KPI",
    "Statistics",
    "Business Analysis",
    "Requirements Gathering",
    "Stakeholder Management",
    "UAT",
    "Linux",
    "ServiceNow",
    "Jira",
    "AWS",
    "Azure",
    "Snowflake",
    "Spark",
    "dbt",
    "MySQL",
    "MariaDB",
    "Oracle",
    "Git",
    "Machine Learning",
]

SKILL_ALIASES = {
    "Power BI": ["power bi", "powerbi"],
    "Power Query": ["power query"],
    "Data Modeling": [
        "data modeling",
        "data modelling",
        "star schema",
        "dimensional model",
    ],
    "Data Analysis": [
        "data analysis",
        "data analytics",
    ],
    "Data Visualization": [
        "data visualization",
        "data visualisation",
    ],
    "Business Analysis": [
        "business analysis",
        "business analyst",
    ],
    "Requirements Gathering": [
        "requirements gathering",
        "requirement gathering",
        "business requirements",
    ],
    "Stakeholder Management": [
        "stakeholder management",
        "stakeholder communication",
        "stakeholders",
    ],
    "Machine Learning": [
        "machine learning",
        "scikit-learn",
        "sklearn",
    ],
}

WEAK_PHRASES = {
    "worked on": "Handled",
    "responsible for": "Managed",
    "helped with": "Supported",
    "involved in": "Contributed to",
    "did": "Completed",
    "made": "Created",
    "used": "Applied",
}

ACTION_VERBS = {
    "analyzed",
    "automated",
    "built",
    "configured",
    "created",
    "designed",
    "developed",
    "diagnosed",
    "implemented",
    "improved",
    "managed",
    "monitored",
    "optimized",
    "prepared",
    "resolved",
    "supported",
    "troubleshot",
    "validated",
}


def _clean(value: Any) -> str:
    if value is None:
        return ""

    text = re.sub(
        r"\s+",
        " ",
        str(value),
    ).strip()

    if text.lower() in {
        "",
        "none",
        "nan",
        "null",
    }:
        return ""

    return text


def _normalise_heading(
    text: str,
) -> str:
    return re.sub(
        r"[^a-z ]+",
        "",
        _clean(text).lower(),
    ).strip()


def _section_name(
    text: str,
) -> str | None:
    normalised = _normalise_heading(
        text
    )

    for section, aliases in (
        SECTION_ALIASES.items()
    ):
        if normalised in aliases:
            return section

    return None


def _all_paragraphs(
    document: Document,
):
    for paragraph in document.paragraphs:
        yield paragraph

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def extract_resume_text(
    resume_path: str | Path,
) -> str:
    document = Document(
        str(resume_path)
    )

    return "\n".join(
        paragraph.text.strip()
        for paragraph in _all_paragraphs(
            document
        )
        if paragraph.text.strip()
    )


def extract_sections(
    resume_path: str | Path,
) -> dict[str, list[str]]:
    document = Document(
        str(resume_path)
    )
    sections: dict[
        str,
        list[str],
    ] = {
        "header": [],
        "objective": [],
        "skills": [],
        "experience": [],
        "projects": [],
        "education": [],
        "certifications": [],
        "other": [],
    }

    current = "header"

    for paragraph in _all_paragraphs(
        document
    ):
        text = _clean(
            paragraph.text
        )

        if not text:
            continue

        detected = _section_name(
            text
        )

        if detected:
            current = detected
            continue

        sections.setdefault(
            current,
            [],
        ).append(text)

    return sections


def extract_job_skills(
    job_description: str,
) -> list[str]:
    text = _clean(
        job_description
    ).lower()
    found: list[str] = []

    for skill in KNOWN_SKILLS:
        aliases = SKILL_ALIASES.get(
            skill,
            [skill.lower()],
        )

        if any(
            alias in text
            for alias in aliases
        ):
            found.append(skill)

    return found


def extract_resume_skills(
    resume_text: str,
) -> list[str]:
    text = _clean(
        resume_text
    ).lower()
    found: list[str] = []

    for skill in KNOWN_SKILLS:
        aliases = SKILL_ALIASES.get(
            skill,
            [skill.lower()],
        )

        if any(
            alias in text
            for alias in aliases
        ):
            found.append(skill)

    return found


def _bullet_score(
    text: str,
) -> dict[str, Any]:
    cleaned = _clean(text)
    lower = cleaned.lower()
    words = cleaned.split()

    starts_with_action = bool(
        words
        and words[0]
        .lower()
        .rstrip(",:;-")
        in ACTION_VERBS
    )

    contains_metric = bool(
        re.search(
            r"\b\d+(?:\.\d+)?%?\b",
            cleaned,
        )
    )

    weak_phrases = [
        phrase
        for phrase in WEAK_PHRASES
        if phrase in lower
    ]

    score = 40

    if starts_with_action:
        score += 25

    if contains_metric:
        score += 20

    if 8 <= len(words) <= 32:
        score += 15

    score -= min(
        len(weak_phrases) * 15,
        30,
    )

    return {
        "text": cleaned,
        "score": max(
            0,
            min(score, 100),
        ),
        "starts_with_action": (
            starts_with_action
        ),
        "contains_metric": (
            contains_metric
        ),
        "weak_phrases": (
            weak_phrases
        ),
        "word_count": len(words),
    }


def _safe_rewrite(
    text: str,
) -> str:
    rewritten = _clean(text)

    for weak, replacement in (
        WEAK_PHRASES.items()
    ):
        pattern = re.compile(
            rf"\b{re.escape(weak)}\b",
            flags=re.IGNORECASE,
        )

        if pattern.search(
            rewritten
        ):
            rewritten = pattern.sub(
                replacement,
                rewritten,
                count=1,
            )
            break

    if rewritten:
        rewritten = (
            rewritten[0].upper()
            + rewritten[1:]
        )

    return rewritten


def review_resume(
    *,
    resume_path: str | Path,
    job_description: str,
    profile_skills: list[str] | None = None,
) -> dict[str, Any]:
    path = Path(resume_path)

    if not path.exists():
        raise FileNotFoundError(
            f"Resume was not found: {path}"
        )

    resume_text = extract_resume_text(
        path
    )
    sections = extract_sections(
        path
    )
    jd_skills = extract_job_skills(
        job_description
    )
    resume_skills = extract_resume_skills(
        resume_text
    )

    owned_skills = {
        _clean(skill)
        for skill in (
            profile_skills or []
        )
        if _clean(skill)
    }
    owned_skills.update(
        resume_skills
    )

    matched = [
        skill
        for skill in jd_skills
        if skill in owned_skills
    ]
    missing = [
        skill
        for skill in jd_skills
        if skill not in owned_skills
    ]

    candidate_bullets = (
        sections["experience"]
        + sections["projects"]
    )

    bullet_reviews = [
        _bullet_score(text)
        for text in candidate_bullets
        if len(text.split()) >= 5
    ]

    weak_bullets = [
        review
        for review in bullet_reviews
        if review["score"] < 70
    ]

    section_scores = {
        "objective": (
            100
            if sections["objective"]
            else 0
        ),
        "skills": (
            min(
                100,
                len(resume_skills) * 8,
            )
        ),
        "experience": (
            min(
                100,
                len(
                    sections["experience"]
                )
                * 18,
            )
        ),
        "projects": (
            min(
                100,
                len(
                    sections["projects"]
                )
                * 20,
            )
        ),
        "education": (
            100
            if sections["education"]
            else 0
        ),
    }

    skill_coverage = (
        round(
            len(matched)
            / len(jd_skills)
            * 100,
        )
        if jd_skills
        else 0
    )

    bullet_quality = (
        round(
            sum(
                item["score"]
                for item in bullet_reviews
            )
            / len(bullet_reviews)
        )
        if bullet_reviews
        else 0
    )

    completeness = round(
        sum(
            section_scores.values()
        )
        / len(section_scores)
    )

    overall_score = round(
        skill_coverage * 0.5
        + bullet_quality * 0.3
        + completeness * 0.2
    )

    recommendations: list[str] = []

    if missing:
        recommendations.append(
            "Do not add unsupported skills: "
            + ", ".join(missing[:8])
            + ". Learn them or leave them as gaps."
        )

    if matched:
        recommendations.append(
            "Emphasize these verified skills near the top: "
            + ", ".join(matched[:8])
            + "."
        )

    if weak_bullets:
        recommendations.append(
            f"Strengthen {len(weak_bullets)} experience/project bullet(s) "
            "with clearer action verbs and outcomes."
        )

    if not any(
        item["contains_metric"]
        for item in bullet_reviews
    ):
        recommendations.append(
            "Add genuine numbers where available, such as ticket volume, "
            "time saved, report frequency, users supported, or dashboard count."
        )

    if not sections["projects"]:
        recommendations.append(
            "Add the most relevant analytics project for this job."
        )

    return {
        "overall_score": overall_score,
        "skill_coverage": skill_coverage,
        "bullet_quality": bullet_quality,
        "completeness": completeness,
        "job_skills": jd_skills,
        "resume_skills": resume_skills,
        "matched_skills": matched,
        "missing_skills": missing,
        "section_scores": section_scores,
        "bullet_reviews": bullet_reviews,
        "weak_bullets": weak_bullets,
        "recommendations": recommendations,
        "sections": sections,
    }


def _build_objective(
    *,
    role: str,
    matched_skills: list[str],
) -> str:
    role_text = (
        _clean(role)
        or "Data Analyst"
    )
    skill_text = ", ".join(
        matched_skills[:6]
    )

    if skill_text:
        skill_clause = (
            f"with practical experience in {skill_text}"
        )
    else:
        skill_clause = (
            "with experience in analytics, reporting, "
            "and enterprise support"
        )

    return (
        f"Detail-oriented {role_text} candidate with 1.1 years of "
        f"enterprise support experience and {skill_clause}. "
        "Experienced in troubleshooting live systems, preparing reports, "
        "building analytical solutions, and supporting data-driven decisions."
    )


def _replace_objective(
    document: Document,
    objective: str,
) -> bool:
    paragraphs = list(
        _all_paragraphs(document)
    )

    for index, paragraph in enumerate(
        paragraphs
    ):
        if _section_name(
            paragraph.text
        ) != "objective":
            continue

        for next_paragraph in (
            paragraphs[index + 1 :]
        ):
            if _section_name(
                next_paragraph.text
            ):
                break

            if _clean(
                next_paragraph.text
            ):
                next_paragraph.text = (
                    objective
                )
                return True

    return False


def _reorder_skills_paragraph(
    document: Document,
    matched_skills: list[str],
) -> bool:
    paragraphs = list(
        _all_paragraphs(document)
    )

    for index, paragraph in enumerate(
        paragraphs
    ):
        if _section_name(
            paragraph.text
        ) != "skills":
            continue

        for next_paragraph in (
            paragraphs[index + 1 :]
        ):
            if _section_name(
                next_paragraph.text
            ):
                break

            original = _clean(
                next_paragraph.text
            )

            if not original:
                continue

            separators = re.split(
                r"[,|•;/]+",
                original,
            )
            existing = [
                _clean(item)
                for item in separators
                if _clean(item)
            ]

            if len(existing) < 2:
                continue

            ordered: list[str] = []

            for matched in matched_skills:
                for skill in existing:
                    if (
                        matched.lower()
                        in skill.lower()
                        and skill
                        not in ordered
                    ):
                        ordered.append(
                            skill
                        )

            for skill in existing:
                if skill not in ordered:
                    ordered.append(skill)

            next_paragraph.text = (
                ", ".join(ordered)
            )
            return True

    return False


def _improve_weak_bullets(
    document: Document,
) -> int:
    updated = 0

    for paragraph in _all_paragraphs(
        document
    ):
        text = _clean(
            paragraph.text
        )

        if len(text.split()) < 5:
            continue

        review = _bullet_score(
            text
        )

        if (
            review["weak_phrases"]
            and review["score"] < 70
        ):
            rewritten = _safe_rewrite(
                text
            )

            if rewritten != text:
                paragraph.text = rewritten
                updated += 1

    return updated


def generate_tailored_resume(
    *,
    source_resume_path: str | Path,
    output_path: str | Path,
    job_description: str,
    role: str,
    profile_skills: list[str] | None = None,
) -> dict[str, Any]:
    review = review_resume(
        resume_path=source_resume_path,
        job_description=job_description,
        profile_skills=profile_skills,
    )

    document = Document(
        str(source_resume_path)
    )

    objective = _build_objective(
        role=role,
        matched_skills=review[
            "matched_skills"
        ],
    )

    objective_updated = (
        _replace_objective(
            document,
            objective,
        )
    )
    skills_reordered = (
        _reorder_skills_paragraph(
            document,
            review[
                "matched_skills"
            ],
        )
    )
    bullets_improved = (
        _improve_weak_bullets(
            document
        )
    )

    destination = Path(
        output_path
    )
    destination.parent.mkdir(
        parents=True,
        exist_ok=True,
    )
    document.save(
        str(destination)
    )

    return {
        "output_path": str(
            destination
        ),
        "review": review,
        "objective": objective,
        "objective_updated": (
            objective_updated
        ),
        "skills_reordered": (
            skills_reordered
        ),
        "bullets_improved": (
            bullets_improved
        ),
    }
